"""
Quick script to:
1. Create the resume_analyses table in the Neon DB (if not exists)
2. Generate a sample DOCX resume
3. POST it to the /api/v1/resume/analyze endpoint
"""
import asyncio
import sys
import os

# ── 1. Ensure DB table exists ──────────────────────────────────────
async def ensure_table():
    import asyncpg
    from dotenv import load_dotenv
    load_dotenv(os.path.join(os.path.dirname(__file__), ".env"))
    
    db_url = os.environ["DATABASE_URL"]
    
    conn = await asyncpg.connect(db_url)
    await conn.execute("""
        CREATE TABLE IF NOT EXISTS resume_analyses (
            id              UUID PRIMARY KEY,
            session_id      UUID NOT NULL UNIQUE,
            file_name       VARCHAR(255) NOT NULL,
            target_role     VARCHAR(255) NOT NULL,
            resume_text     TEXT NOT NULL,
            overall_score   INTEGER NOT NULL,
            grade           VARCHAR(2) NOT NULL,
            ats_score       INTEGER NOT NULL,
            estimated_level VARCHAR(10) NOT NULL,
            top_strengths       JSONB NOT NULL DEFAULT '[]'::jsonb,
            critical_weaknesses JSONB NOT NULL DEFAULT '[]'::jsonb,
            keywords_missing    JSONB NOT NULL DEFAULT '[]'::jsonb,
            sections            JSONB NOT NULL DEFAULT '[]'::jsonb,
            created_at      TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP
        );
        CREATE INDEX IF NOT EXISTS ix_resume_analyses_session_id 
            ON resume_analyses (session_id);
    """)
    print("[OK] resume_analyses table ready")
    await conn.close()

# ── 2. Create a sample DOCX resume ─────────────────────────────────
def create_sample_resume() -> str:
    from docx import Document
    doc = Document()
    doc.add_heading("Ribhu Sharma", level=1)
    doc.add_paragraph("Senior Software Engineer | Python | FastAPI | AI/ML")
    doc.add_heading("Summary", level=2)
    doc.add_paragraph(
        "Experienced software engineer with 5+ years building scalable "
        "backend systems using Python, FastAPI, and PostgreSQL. Passionate "
        "about AI/ML engineering and building intelligent applications."
    )
    doc.add_heading("Work Experience", level=2)
    doc.add_paragraph(
        "Senior Backend Engineer — TechCorp (2022-Present)\n"
        "- Designed and built microservices handling 50K+ RPM using FastAPI\n"
        "- Implemented async data pipelines with SQLAlchemy 2.0 and asyncpg\n"
        "- Integrated Claude API for intelligent document processing\n"
        "- Reduced API latency by 40% through Redis caching strategies"
    )
    doc.add_paragraph(
        "Software Engineer — DataFlow Inc (2019-2022)\n"
        "- Built REST APIs serving 10M+ requests/day with Python/Flask\n"
        "- Designed PostgreSQL schemas for time-series analytics\n"
        "- Implemented CI/CD pipelines with GitHub Actions and Docker"
    )
    doc.add_heading("Skills", level=2)
    doc.add_paragraph(
        "Python, FastAPI, Flask, SQLAlchemy, PostgreSQL, Redis, Docker, "
        "Kubernetes, AWS, Claude API, LangChain, PyTorch, TensorFlow, "
        "Git, CI/CD, REST APIs, GraphQL, Microservices"
    )
    doc.add_heading("Education", level=2)
    doc.add_paragraph("B.Tech Computer Science — IIT Delhi (2015-2019)")
    doc.add_heading("Projects", level=2)
    doc.add_paragraph(
        "CareerOS — AI-powered career operating system\n"
        "- Built resume analysis engine using Claude API\n"
        "- Implemented event-sourced progress tracking\n"
        "- Tech: FastAPI, PostgreSQL, Redis, Anthropic SDK"
    )

    path = os.path.join(os.path.dirname(__file__), "sample_resume.docx")
    doc.save(path)
    print(f"[OK] Sample resume saved to {path}")
    return path

# ── 3. Hit the API ──────────────────────────────────────────────────
def run_analyze(resume_path: str):
    import httpx

    api_base = os.environ.get("API_BASE_URL", "http://127.0.0.1:8000").rstrip("/")
    url = f"{api_base}/api/v1/resume/analyze"

    with open(resume_path, "rb") as f:
        files = {"file": ("sample_resume.docx", f, 
                 "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        data = {"target_role": "Senior AI/ML Engineer"}

        print(f"\n[->] POST {url}")
        print(f"    file: sample_resume.docx")
        print(f"    target_role: Senior AI/ML Engineer")
        print(f"    Waiting for Claude analysis (this may take 10-30 seconds)...\n")

        resp = httpx.post(url, files=files, data=data, timeout=120.0)

    print(f"[<-] Status: {resp.status_code}")
    
    if resp.status_code == 200:
        import json
        result = resp.json()
        print(f"\n{'='*60}")
        print(f"  RESUME ANALYSIS RESULT")
        print(f"{'='*60}")
        print(f"  Overall Score : {result['overall_score']}/100")
        print(f"  Grade         : {result['grade']}")
        print(f"  ATS Score     : {result['ats_score']}/100")
        print(f"  Level         : {result['estimated_level']}")
        print(f"{'='*60}")
        print(f"\n  Top Strengths:")
        for s in result['top_strengths']:
            print(f"    + {s}")
        print(f"\n  Critical Weaknesses:")
        for w in result['critical_weaknesses']:
            print(f"    - {w}")
        print(f"\n  Missing Keywords:")
        for k in result['keywords_missing']:
            print(f"    * {k}")
        print(f"\n  Section Scores:")
        for sec in result['sections']:
            status_icon = {"strong": "[STRONG]", "needs_improvement": "[IMPROVE]", "weak": "[WEAK]", "missing": "[MISSING]"}.get(sec['status'], "?")
            print(f"    {status_icon} {sec['name']}: {sec['score']}/100")
            print(f"      {sec['feedback']}")
        print(f"\n{'='*60}")
        print(f"  Session ID: {result['session_id']}")
        print(f"  Record ID:  {result['id']}")
        print(f"{'='*60}")
    else:
        print(f"  Error: {resp.text}")

# ── Main ────────────────────────────────────────────────────────────
if __name__ == "__main__":
    print("=" * 60)
    print("  CareerOS Resume Analyzer — End-to-End Test")
    print("=" * 60)
    
    # Step 1: Ensure table
    asyncio.run(ensure_table())
    
    # Step 2: Create sample resume
    resume_path = create_sample_resume()
    
    # Step 3: Test the API
    run_analyze(resume_path)
